"""results_table.py — lưu kết quả từng lần chạy ra JSON, rồi điền experiments.xlsx từ mẫu
templates/experiment_table_template.xlsx.

Tên cột của sheet "Experiments" (giữ nguyên, đúng thứ tự mẫu):
    exp_id, group, description, loss, optimizer, lr, weight_decay, batch, epochs, hidden, dropout,
    clip_norm, precision, init, seed, step0_loss, best_val_loss, best_epoch, final_train_loss,
    final_val_loss, val_acc, val_macro_f1, time_per_epoch_s, peak_mem_MB, diverged,
    eval_acc, eval_macro_f1, figure_file, notes
(các cột công thức ở cuối bảng mẫu tự tính, không ghi đè)
"""
from __future__ import annotations

import json
import math
from pathlib import Path

FORMULA_COLS = ("step0_gap_vs_lnC", "gap_val_minus_train", "delta_val_f1_vs_base", "beyond_noise")
MAX_ROWS = 60  # công thức của mẫu phủ dòng 2..61
OPT_NAMES = {"sgd": "SGD", "sgd_momentum": "SGD+momentum", "adam": "Adam", "adamw": "AdamW"}


def _clean(x):
    """JSON không có NaN/inf chuẩn -> đổi thành None; tuple -> list."""
    if isinstance(x, float) and not math.isfinite(x):
        return None
    if isinstance(x, dict):
        return {k: _clean(v) for k, v in x.items()}
    if isinstance(x, (list, tuple)):
        return [_clean(v) for v in x]
    return x


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi cfg/history/summary (KHÔNG ghi best_state) ra <results_dir>/<exp_id>.json."""
    Path(results_dir).mkdir(parents=True, exist_ok=True)
    path = Path(results_dir) / f"{result['cfg']['exp_id']}.json"
    payload = {k: _clean(result[k]) for k in ("cfg", "history", "summary")}
    for k in ("notes", "eval"):
        if k in result:
            payload[k] = _clean(result[k])
    with open(path, "w", encoding="utf-8") as f:
        json.dump(payload, f, indent=2, ensure_ascii=False)
    return str(path)


def load_result(path: str) -> dict:
    with open(path, encoding="utf-8") as f:
        r = json.load(f)
    nan = float("nan")
    r["summary"] = {k: (nan if v is None and k != "best_epoch" else v) for k, v in r["summary"].items()}
    r["history"] = {k: [nan if v is None else v for v in vs] for k, vs in r["history"].items()}
    return r


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi *.json trong results_dir, sắp theo exp_id."""
    return sorted((load_result(p) for p in Path(results_dir).glob("*.json")),
                  key=lambda r: r["cfg"]["exp_id"])


def _num(x, nd: int | None = None):
    if x is None or (isinstance(x, float) and not math.isfinite(x)):
        return None
    return round(x, nd) if nd is not None and isinstance(x, float) else x


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """cfg + summary (+ eval_acc/eval_macro_f1 nếu có) -> một dòng bảng; khoá trùng tên cột."""
    c, s = result["cfg"], result["summary"]
    row = dict(
        exp_id=c["exp_id"], group=c["group"], description=c["description"],
        loss=c["loss"].upper(), optimizer=OPT_NAMES[c["optimizer"]], lr=c["lr"],
        weight_decay=c["weight_decay"], batch=c["batch"], epochs=c["epochs"],
        hidden="-".join(str(h) for h in c["hidden"]), dropout=c["dropout"],
        clip_norm="none" if c["clip_norm"] is None else c["clip_norm"],
        precision=c["precision"], init=c["init"], seed=c["seed"],
        step0_loss=_num(s["step0_loss"], 4), best_val_loss=_num(s["best_val_loss"], 4),
        best_epoch=s["best_epoch"], final_train_loss=_num(s["final_train_loss"], 4),
        final_val_loss=_num(s["final_val_loss"], 4), val_acc=_num(s["val_acc"], 4),
        val_macro_f1=_num(s["val_macro_f1"], 4), time_per_epoch_s=_num(s["time_per_epoch_s"], 2),
        peak_mem_MB=_num(s["peak_mem_MB"], 1), diverged="Y" if s["diverged"] else "N",
        eval_acc=None, eval_macro_f1=None,
        figure_file=f"figures/{c['exp_id']}.png", notes=notes or result.get("notes", ""),
    )
    if eval_scores is not None:
        row["eval_acc"] = round(eval_scores["accuracy"], 4)
        row["eval_macro_f1"] = round(eval_scores["macro_f1"], 4)
    return row


def write_xlsx(rows: list[dict], template_path: str, out_path: str,
               seed_ids: list[str] | None = None, summary_notes: dict | None = None) -> None:
    """Điền sheet Experiments (từ dòng 2), sheet Seeds (cột A) và cột nhận xét của Summary, rồi lưu."""
    import openpyxl

    if len(rows) > MAX_ROWS:
        raise ValueError(f"mẫu chỉ có công thức cho {MAX_ROWS} dòng, đang có {len(rows)}")
    wb = openpyxl.load_workbook(template_path)  # KHÔNG data_only=True (sẽ mất công thức)
    ws = wb["Experiments"]
    header = [cell.value for cell in ws[1]]
    for i, row in enumerate(rows, start=2):
        for j, col in enumerate(header, start=1):
            if col is None or col in FORMULA_COLS:
                continue
            ws.cell(row=i, column=j, value=row.get(col))
    for i in range(len(rows) + 2, MAX_ROWS + 2):  # xoá giá trị mẫu còn sót (dòng baseline ví dụ)
        for j, col in enumerate(header, start=1):
            if col is not None and col not in FORMULA_COLS:
                ws.cell(row=i, column=j, value=None)

    if seed_ids is not None:
        wsS = wb["Seeds"]
        for k in range(5):  # Seeds!A2:A6
            wsS.cell(row=2 + k, column=1, value=seed_ids[k] if k < len(seed_ids) else None)

    if summary_notes:
        wsM = wb["Summary"]
        for r in range(2, wsM.max_row + 1):
            g = wsM.cell(row=r, column=1).value
            if g in summary_notes:
                wsM.cell(row=r, column=8, value=summary_notes[g])
    wb.save(out_path)
